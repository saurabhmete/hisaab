"""Import data from the 'Salary Plan' style Excel workbook.

Layout expectations:
- Year sheets named like '2025': row 1 = Total Salary, row 2 = month names,
  following rows = expense categories, until the 'Total Expense' row.
  A 'Can Invest' row holds the invested amount. Derived rows are skipped.
- An 'Investment Tracker' sheet: row 1 = month names, then labeled rows for
  India/Germany values and profits and the EUR->INR rate. Derived rows are
  recomputed by the app, not imported.
"""
import re
from openpyxl import load_workbook

from . import database as d

SKIP_ROWS = (
    "total expense", "salary remaining", "remainder in main account",
    "expenses", "total salary", "can invest",
)

INV_ROWS = {
    "india investment value": "india_value",
    "india profit (": "india_profit",
    "germany investment value": "germany_value",
    "germany profit (": "germany_profit",
    "eur to inr rate": "eur_inr_rate",
}


def _month_no(name) -> int | None:
    if not isinstance(name, str):
        return None
    name = name.strip().capitalize()
    return d.MONTH_NAMES.index(name) + 1 if name in d.MONTH_NAMES else None


def import_workbook(db, path_or_file, investment_year: int | None = None) -> dict:
    """Import all recognizable sheets. Returns counts per sheet."""
    wb = load_workbook(path_or_file, data_only=True)
    report = {}
    year_sheets = [ws for ws in wb.worksheets if re.fullmatch(r"\d{4}", ws.title)]
    for ws in year_sheets:
        report[ws.title] = _import_year_sheet(db, ws, int(ws.title))

    inv = next((ws for ws in wb.worksheets if "investment" in ws.title.lower()), None)
    if inv is not None:
        # The tracker sheet has no year of its own; default to the latest year sheet.
        year = investment_year or (max(int(ws.title) for ws in year_sheets) if year_sheets else None)
        if year:
            report[inv.title] = _import_investment_sheet(db, inv, year)
    return report


def _import_year_sheet(db, ws, year: int) -> int:
    rows = list(ws.iter_rows(values_only=False))
    if len(rows) < 2:
        return 0
    header = rows[1]  # month names
    months = {}  # column index -> month_id
    prev_m, y = 0, year
    for idx, cell in enumerate(header[1:], start=1):
        m = _month_no(cell.value)
        if m is None:
            continue
        if m < prev_m:  # wrapped past December: e.g. a 2025 sheet ending in January 2026
            y += 1
        prev_m = m
        salary = rows[0][idx].value
        salary = float(salary) if isinstance(salary, (int, float)) else 0.0
        months[idx] = d.upsert_month(db, y, m, salary=salary)

    count = 0
    for row in rows[2:]:
        # Label is normally in column A, but tolerate rows where it landed in
        # another cell (e.g. a 'Furniture' label typed mid-row).
        label = next(
            (c.value for c in row if isinstance(c.value, str) and c.value.strip()), None
        )
        if label is None:
            continue
        low = label.strip().lower()
        if low.startswith("can invest"):
            for idx, month_id in months.items():
                v = row[idx].value if idx < len(row) else None
                if isinstance(v, (int, float)):
                    db.execute("UPDATE months SET invested=? WHERE id=?", (float(v), month_id))
            continue
        if any(low.startswith(s) for s in SKIP_ROWS):
            continue
        for idx, month_id in months.items():
            v = row[idx].value if idx < len(row) else None
            if isinstance(v, (int, float)):
                d.set_expense(db, month_id, label.strip(), float(v))
                count += 1
    return count


def _import_investment_sheet(db, ws, year: int) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    header = rows[0]
    month_cols = {i: _month_no(v) for i, v in enumerate(header) if _month_no(v)}

    data = {}  # month -> field -> value
    for row in rows[1:]:
        label = row[0]
        if not isinstance(label, str):
            continue
        low = label.strip().lower()
        field = next((f for key, f in INV_ROWS.items() if key in low), None)
        if field is None:
            continue
        for idx, m in month_cols.items():
            v = row[idx] if idx < len(row) else None
            if isinstance(v, (int, float)):
                data.setdefault(m, {})[field] = float(v)

    for m, fields in sorted(data.items()):
        d.upsert_investment(
            db, year, m,
            fields.get("india_value"), fields.get("india_profit"),
            fields.get("germany_value"), fields.get("germany_profit"),
            fields.get("eur_inr_rate", 110.0),
        )
    return len(data)
